#!/usr/bin/env python3
"""
grab.py v4 — универсальный загрузчик клипов с YouTube под монтаж (Premiere / DaVinci).

Установка:  pip install -U "yt-dlp[default]" openpyxl ;  ffmpeg + ffprobe в PATH ;  deno (winget install DenoLand.Deno)

ВХОД (первый аргумент):
  *.xlsx   таблица; колонки находятся по названию (см. ALIASES): запрос, таймкод, статус, блок, момент, №.
           Нужна только колонка запроса. Скрипт сам пишет в неё же статус/ссылку/файл (бэкап: *.bak.xlsx).
  *.csv    то же, что xlsx, но без записи результатов в файл (только log.csv).
  *.txt    одна строка = один клип:   запрос или ссылка | таймкод | название        (# — комментарий)
  иное     считается одиночным запросом/ссылкой:   python grab.py "https://youtu.be/XXXX" --tc 1:23-1:40

ТАЙМКОДЫ: 1:23-1:40 | 01:02:03-01:02:30 | 85-110 (секунды) | 0:45+15s (старт+длительность) |
          1:23- (до конца) | -0:40 (с начала) | 1:23 (одна точка: --clip-len секунд от неё) |
          ячейки Excel типа «время» (см. ниже). Разделитель: - – — .. > to до.
          Если Excel превратил «1:23» во время 01:23:00, оно читается как 1 мин 23 с (секунды=0).
          Надёжнее: формат колонки «Текст».

ПРИМЕРЫ:
  python grab.py list.xlsx --dry-run                    # только поиск и фильтры
  python grab.py list.xlsx --priority A --pick          # выбирать кандидата вручную
  python grab.py list.xlsx --only 5,9-14 --verbose
  python grab.py clips.txt --out D:\\clips --proxy socks5://127.0.0.1:1080
Выход: <out>/<Блок>/<NN>_<название>.mp4 — H.264 + AAC, yuv420p, faststart. Лог запуска: <out>/run_*.txt
"""
import sys
for _s in (sys.stdout, sys.stderr):
    if _s and hasattr(_s, 'reconfigure'):
        try: _s.reconfigure(encoding='utf-8', errors='replace')
        except Exception: pass
import argparse, csv, datetime as dt, json, platform, random, re, shutil, subprocess, sys, time, traceback
from pathlib import Path

try:
    import yt_dlp
    import yt_dlp.version
    from yt_dlp.utils import download_range_func
except ImportError:
    sys.exit('Нужен yt-dlp:  pip install -U "yt-dlp[default]"')

ALIASES = {
    "n": ["№", "n", "#", "номер", "id"],
    "block": ["блок", "block", "группа", "раздел", "group"],
    "title": ["момент сценария", "момент", "название", "title", "описание", "name"],
    "query": ["поисковый запрос", "запрос", "query", "url", "link"],
    "tc": ["таймкод (от-до)", "таймкод", "timecode", "tc"],
    "status": ["статус", "status"],
    "prio": ["приоритет", "priority", "prio"],
}
OUT_COLS = {"status": "Статус", "url": "Ссылка (выбрано)", "file": "Файл", "note": "Заметка"}
PENDING = {"", "найти", "find", "todo", "pending"}


class SkipItem(Exception):
    """Позицию не нужно повторять (ошибка данных, а не сети)."""


# ───────────────────────── вывод / лог ─────────────────────────
class Tee:
    def __init__(self, stream, f):
        self.stream, self.f = stream, f
    def write(self, d):
        try:
            self.stream.write(d)
        except UnicodeEncodeError:
            self.stream.write(d.encode("ascii", "replace").decode())
        try:
            self.f.write(d); self.f.flush()
        except Exception:
            pass
    def flush(self):
        self.stream.flush()
    def isatty(self):
        return False


class YLog:
    def __init__(self, verbose):
        self.verbose = verbose
    def debug(self, m):
        if m.startswith("[debug]") and not self.verbose:
            return
        print(m)
    def info(self, m): print(m)
    def warning(self, m): print(f"WARNING: {m}")
    def error(self, m): print(f"ERROR: {m}")


def slug(s, n=40):
    s = re.sub(r"[^\w\-]+", "_", str(s or ""), flags=re.U).strip("_")
    return s[:n] or "clip"


def fmt_dur(d):
    if not d:
        return "?"
    d = int(d)
    return f"{d//3600}:{d%3600//60:02d}:{d%60:02d}" if d >= 3600 else f"{d//60}:{d%60:02d}"


def make_hook(label):
    last = [0.0]
    def hook(d):
        out = sys.__stdout__
        if d["status"] == "downloading":
            now = time.time()
            if now - last[0] < 0.5:
                return
            last[0] = now
            tot = d.get("total_bytes") or d.get("total_bytes_estimate")
            done = d.get("downloaded_bytes") or 0
            pct = f"{done / tot * 100:5.1f}%" if tot else "   ?  "
            sp = d.get("speed")
            eta = d.get("eta")
            line = f"\r  {label} {pct}  {sp / 1e6:.1f} MB/s  ETA {fmt_dur(eta) if eta is not None else '?'}" if sp else f"\r  {label} {pct}"
            out.write(line.ljust(70)); out.flush()
        elif d["status"] == "finished":
            out.write("\r" + " " * 70 + "\r"); out.flush()
    return hook


# ───────────────────────── таймкоды ─────────────────────────
SEP = re.compile(r"\s*(?:\.\.|[-–—>->]|\bto\b|\bдо\b)\s*", re.I)


def parse_point(s):
    s = s.strip().lower().replace(",", ".")
    if not s:
        return None
    if re.fullmatch(r"\d+(?::\d+){1,2}(?:\.\d+)?", s):
        v = 0.0
        for p in s.split(":"):
            v = v * 60 + float(p)
        return v
    m = re.fullmatch(r"(?:(\d+)\s*(?:h|ч)\s*)?(?:(\d+)\s*(?:min|мин|m|м)\s*)?(?:(\d+(?:\.\d+)?)\s*(?:sec|сек|s|с)?)?", s)
    if m and any(m.groups()):
        h, mi, se = m.groups()
        return int(h or 0) * 3600 + int(mi or 0) * 60 + float(se or 0)
    raise ValueError(f"не разобрана точка времени: {s!r}")


def _clock_fix(h, m, s):
    """Excel превращает «1:23» во время 01:23:00 -> считаем как 1 мин 23 с."""
    return (h * 60 + m) if s == 0 and h < 24 else (h * 3600 + m * 60 + s)


def parse_tc(v, clip_len=15.0):
    """-> (start, end|None) или None. end=None: до конца ролика."""
    if v is None or (isinstance(v, str) and v.strip() in ("", "—", "-")):
        return None
    if isinstance(v, dt.datetime):
        v = v.time()
    if isinstance(v, dt.time):
        st = _clock_fix(v.hour, v.minute, v.second + v.microsecond / 1e6)
        return st, st + clip_len
    if isinstance(v, dt.timedelta):
        t = v.total_seconds()
        st = _clock_fix(int(t // 3600), int(t % 3600 // 60), t % 60)
        return st, st + clip_len
    if isinstance(v, (int, float)):
        return float(v), float(v) + clip_len
    s = str(v).strip()
    if "+" in s:
        a, b = s.split("+", 1)
        st = parse_point(a)
        if st is None:
            raise ValueError(f"нет начала: {s!r}")
        return st, st + parse_point(b)
    parts = SEP.split(s, maxsplit=1)
    if len(parts) == 1:
        st = parse_point(parts[0])
        return st, st + clip_len
    a, b = (parse_point(parts[0]) or 0.0), parse_point(parts[1])
    if b is not None and b <= a:
        raise ValueError(f"конец раньше начала: {s!r}")
    return a, b


def parse_only(spec):
    out = set()
    for part in spec.split(","):
        if "-" in part:
            a, b = part.split("-")
            out.update(range(int(a), int(b) + 1))
        elif part.strip():
            out.add(int(part))
    return out


# ───────────────────────── ввод ─────────────────────────
def norm(x):
    return str(x or "").strip().lower()


def colmap_from(headers):
    h = [norm(x) for x in headers]
    return {k: next((i for i, name in enumerate(h) if name in al), None) for k, al in ALIASES.items()}


def mk_item(n, block, title, query, tc, status, prio, row=None):
    try:
        n = int(float(n))
    except Exception:
        n = None
    return {"n": n, "block": str(block or ""), "title": str(title or ""), "query": str(query or "").strip(),
            "tc": tc, "status": str(status or "").strip(), "prio": str(prio or "").strip(), "row": row}


class XlsxSync:
    FILL = {"ok": "C6EFCE", "err": "F8CBAD"}

    def __init__(self, path, sheet):
        from openpyxl import load_workbook
        self.path = Path(path)
        self.wb = load_workbook(self.path)
        names = self.wb.sheetnames
        self.ws = self.wb[sheet] if sheet in names else (self.wb["Монтажный лист"] if "Монтажный лист" in names else self.wb.worksheets[0])
        hdr = [c.value for c in self.ws[1]]
        self.cm = colmap_from(hdr)
        if self.cm["query"] is None:
            sys.exit(f"В листе «{self.ws.title}» нет колонки запроса. Заголовки: {hdr}. Ожидается одно из: {ALIASES['query']}")
        bak = self.path.with_name(self.path.stem + ".bak.xlsx")
        if not bak.exists():
            shutil.copy2(self.path, bak)
        self.out = {}
        lower = {norm(c.value): c.column for c in self.ws[1] if c.value}
        for key, name in OUT_COLS.items():
            col = lower.get(norm(name)) or (self.cm["status"] + 1 if key == "status" and self.cm["status"] is not None else None)
            if key == "status" and self.cm["status"] is not None:
                col = self.cm["status"] + 1
            if not col:
                col = self.ws.max_column + 1
                cell = self.ws.cell(1, col, name)
                cell._style = self.ws.cell(1, 1)._style
                self.ws.column_dimensions[cell.column_letter].width = 28 if key != "note" else 40
            self.out[key] = col
        self.warned = False

    def items(self):
        g = lambda r, k: self.ws.cell(r, self.cm[k] + 1).value if self.cm[k] is not None else None
        for r in range(2, self.ws.max_row + 1):
            q = g(r, "query")
            if not q:
                continue
            n = g(r, "n")
            yield mk_item(n if n is not None else r - 1, g(r, "block"), g(r, "title"), q, g(r, "tc"),
                          g(r, "status"), g(r, "prio"), row=r)

    def update(self, item, status=None, url=None, file=None, note=None):
        from openpyxl.styles import PatternFill
        r = item["row"]
        for key, val in (("status", status), ("url", url), ("file", file), ("note", note)):
            if val is not None:
                self.ws.cell(r, self.out[key]).value = val
        if status:
            col = self.FILL["ok"] if status == "скачано" else self.FILL["err"] if status.startswith("ошибка") else None
            if col:
                self.ws.cell(r, self.out["status"]).fill = PatternFill("solid", fgColor=col)
        self.save()

    def save(self):
        try:
            self.wb.save(self.path)
        except PermissionError:
            alt = self.path.with_name(self.path.stem + "_results.xlsx")
            if not self.warned:
                print(f"ВНИМАНИЕ: {self.path.name} открыт в Excel — результаты пишутся в {alt.name}. Закройте Excel, чтобы писать в исходный файл.", file=sys.stderr)
                self.warned = True
            self.wb.save(alt)


def load_items(src, args):
    p = Path(src)
    if p.suffix.lower() == ".xlsx" and p.exists():
        sync = XlsxSync(p, args.sheet)
        return list(sync.items()), sync
    if p.suffix.lower() == ".csv" and p.exists():
        txt = p.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(txt[:2000], delimiters=",;\t")
        rows = list(csv.reader(txt.splitlines(), dialect))
        cm = colmap_from(rows[0])
        if cm["query"] is None:
            sys.exit(f"В csv нет колонки запроса. Ожидается одно из: {ALIASES['query']}")
        g = lambda r, k: r[cm[k]] if cm[k] is not None and cm[k] < len(r) else None
        return [mk_item(g(r, "n") or i, g(r, "block"), g(r, "title"), g(r, "query"), g(r, "tc"), g(r, "status"), g(r, "prio"))
                for i, r in enumerate(rows[1:], 1) if g(r, "query")], None
    if p.suffix.lower() == ".txt" and p.exists():
        items = []
        for line in p.read_text(encoding="utf-8-sig").splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = [x.strip() for x in line.split("|")]
            parts += [""] * (3 - len(parts))
            items.append(mk_item(len(items) + 1, "", parts[2] or parts[0], parts[0], parts[1] or None, "", ""))
        return items, None
    return [mk_item(1, "", args.name or src, src, args.tc, "", "")], None


def is_pending(item, args):
    s = norm(item["status"])
    if args.redo:
        return True
    if args.status:
        return s == norm(args.status)
    return s in PENDING or s.startswith("ошибка") or s.startswith("error")


# ───────────────────────── ffmpeg ─────────────────────────
def probe(path):
    try:
        r = subprocess.run(["ffprobe", "-v", "error", "-print_format", "json", "-show_streams", "-show_format", str(path)],
                           capture_output=True, text=True, timeout=60)
        if r.returncode:
            return None
        j = json.loads(r.stdout)
    except Exception:
        return None
    v = next((s for s in j.get("streams", []) if s.get("codec_type") == "video"), None)
    a = next((s for s in j.get("streams", []) if s.get("codec_type") == "audio"), None)
    try:
        dur = float(j.get("format", {}).get("duration") or (v or {}).get("duration") or 0)
    except Exception:
        dur = 0.0
    return {"v": v and v.get("codec_name"), "a": a and a.get("codec_name"), "pix": v and v.get("pix_fmt"),
            "w": v and v.get("width"), "h": v and v.get("height"), "dur": dur}


def is_valid(path):
    p = Path(path)
    if not p.exists() or p.suffix.lower() != ".mp4" or p.stat().st_size < 10_000:
        return False
    i = probe(p)
    return bool(i and i["v"] and i["dur"] > 0.3)


def encode(src, dst, start=None, end=None, keep_src=False):
    """Приводит к H.264+AAC/yuv420p. Если start/end заданы — точная вырезка с перекодированием."""
    src, dst = Path(src), Path(dst)
    info = probe(src)
    if not info or not info["v"]:
        raise RuntimeError(f"ffprobe: {src.name} не читается")
    ok = info["v"] == "h264" and info["a"] in (None, "aac") and info["pix"] in (None, "yuv420p") and src.suffix.lower() == ".mp4"
    if start is None and end is None and ok:
        (shutil.copy2 if keep_src else shutil.move)(src, dst)
        return
    cmd = ["ffmpeg", "-y", "-hide_banner", "-loglevel", "error"]
    if start:
        cmd += ["-ss", f"{start:.3f}"]
    cmd += ["-i", str(src)]
    if end is not None:
        cmd += ["-t", f"{end - (start or 0):.3f}"]
    cmd += ["-map", "0:v:0", "-map", "0:a:0?", "-c:v", "libx264", "-preset", "veryfast", "-crf", "18", "-pix_fmt", "yuv420p",
            "-c:a", "aac", "-b:a", "192k", "-movflags", "+faststart", str(dst)]
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode:
        raise RuntimeError("ffmpeg: " + r.stderr[-600:])
    if not keep_src:
        src.unlink(missing_ok=True)


# ───────────────────────── yt-dlp ─────────────────────────
def base_opts(args, verbose=False):
    o = {"quiet": True, "noprogress": True, "noplaylist": True, "logger": YLog(verbose), "verbose": verbose,
         "socket_timeout": 30, "retries": 10, "fragment_retries": 10, "extractor_retries": 3, "file_access_retries": 3,
         "retry_sleep_functions": {k: (lambda n: min(2 ** n, 60)) for k in ("http", "fragment", "extractor", "file_access")}}
    if args.cfb:
        o["cookiesfrombrowser"] = (args.cfb,)
    if args.cookies:
        o["cookiefile"] = args.cookies
    if args.proxy:
        o["proxy"] = args.proxy
    return o


def search(query, args):
    opts = {**base_opts(args), "extract_flat": True, "skip_download": True}
    with yt_dlp.YoutubeDL(opts) as y:
        info = y.extract_info(f"ytsearch{args.results * 2}:{query}", download=False)
    res = []
    for e in info.get("entries") or []:
        if e:
            res.append({"url": e.get("url") or f"https://www.youtube.com/watch?v={e['id']}", "title": e.get("title") or "",
                        "dur": e.get("duration"), "chan": e.get("channel") or e.get("uploader") or "",
                        "live": e.get("is_live") or e.get("live_status") in ("is_live", "is_upcoming")})
    return res


def reject_reason(c, has_tc, args):
    if c["live"]:
        return "прямой эфир"
    if "/shorts/" in c["url"]:
        return "Shorts"
    if not c["dur"]:
        return "нет длительности (стрим?)"
    if not has_tc and not args.allow_long and c["dur"] > args.max_min * 60:
        return f"длиннее {args.max_min:g} мин"
    return None


def choose(cands, item, has_tc, args):
    ok, dropped = [], {}
    for c in cands:
        r = reject_reason(c, has_tc, args)
        if r:
            dropped[r] = dropped.get(r, 0) + 1
        else:
            ok.append(c)
    ok = ok[:args.results]
    if dropped:
        print(f"    отсеяно: {dropped}")
    if not ok:
        return None
    if not args.pick:
        return ok[0]
    print(f"\n[{item['n']}] {item['block']} — {item['title']}\n    запрос: {item['query']}")
    for i, c in enumerate(ok, 1):
        print(f"  {i}. {c['title'][:80]}  [{fmt_dur(c['dur'])}]  {c['chan']}\n     {c['url']}")
    a = input("  номер / Enter=1 / s=пропустить / ссылка: ").strip()
    if a.lower() == "s":
        return None
    if a.startswith("http"):
        return {"url": a, "title": "(ссылка вручную)", "dur": None, "chan": "", "live": False}
    return ok[int(a) - 1] if a.isdigit() and 1 <= int(a) <= len(ok) else ok[0]


def fetch_info(url, args):
    with yt_dlp.YoutubeDL({**base_opts(args, args.verbose), "skip_download": True}) as y:
        return y.extract_info(url, download=False)


def fmt_string(h):
    return (f"bv*[vcodec^=avc1][height<={h}]+ba[ext=m4a]/b[vcodec^=avc1][height<={h}]/"
            f"bv*[height<={h}]+ba/b[height<={h}]")


def ydl_download(url, stem, args, label, ranges=None):
    stem = Path(stem)
    stem.parent.mkdir(parents=True, exist_ok=True)
    opts = {**base_opts(args, args.verbose), "format": fmt_string(args.height), "merge_output_format": "mp4",
            "outtmpl": f"{stem}.%(ext)s", "concurrent_fragment_downloads": 4, "progress_hooks": [make_hook(label)]}
    if args.sleep:
        opts["sleep_interval"], opts["max_sleep_interval"] = args.sleep, args.sleep * 2
    if ranges:
        opts["download_ranges"] = download_range_func(None, ranges)
        opts["force_keyframes_at_cuts"] = True
    with yt_dlp.YoutubeDL(opts) as y:
        y.download([url])
    files = [f for f in stem.parent.glob(stem.name + ".*") if f.suffix not in (".part", ".ytdl", ".temp", ".json")]
    if not files:
        raise RuntimeError("yt-dlp не создал файл")
    return max(files, key=lambda f: f.stat().st_size)


# ───────────────────────── обработка одной позиции ─────────────────────────
def process(item, url, tc, args, out, final, label):
    info = fetch_info(url, args)
    if info.get("is_live") or info.get("live_status") in ("is_live", "is_upcoming"):
        raise SkipItem("прямой эфир")
    dur, w, h = info.get("duration"), info.get("width"), info.get("height")
    vid = info.get("id") or slug(url)
    notes = []
    if w and h and h > w:
        notes.append("вертикальное видео (Shorts?)")
    if tc is None and dur and dur > args.max_min * 60 and not args.allow_long:
        raise SkipItem(f"ролик {fmt_dur(dur)} длиннее лимита {args.max_min:g} мин: задайте таймкод или --allow-long")
    s = e = None
    if tc:
        s, e = tc
        if dur and s >= dur:
            raise SkipItem(f"таймкод {fmt_dur(s)} за пределами ролика ({fmt_dur(dur)})")
        if dur and (e is None or e > dur):
            if e is not None:
                notes.append(f"конец таймкода обрезан до длительности ролика {fmt_dur(dur)}")
            e = dur
    mode = args.cut_mode
    if not tc:
        mode = "full"
    elif mode == "auto":
        mode = "section" if (dur and dur > args.full_max_min * 60) else "full"

    tmp = out / "_tmp" / f"{item['n']}"
    shutil.rmtree(tmp, ignore_errors=True); tmp.mkdir(parents=True)
    tmp_final = tmp / "final.mp4"
    try:
        if mode == "full":
            cache = getattr(args, "resolved_cache", out / ".cache"); cache.mkdir(exist_ok=True)
            cfile = cache / f"{vid}.mp4"
            if not is_valid(cfile):
                cfile.unlink(missing_ok=True)
                raw = ydl_download(url, cache / f"{vid}_raw", args, label)
                encode(raw, cfile)
            if tc:
                encode(cfile, tmp_final, s, e, keep_src=True)
            else:
                shutil.copy2(cfile, tmp_final)
        else:
            rs = max(0.0, s - args.margin)
            re_ = (e + args.margin) if e else float("inf")
            raw = ydl_download(url, tmp / "sec", args, label, ranges=[(rs, re_)])
            encode(raw, tmp_final, s - rs, (e - rs) if e else None)
        if not is_valid(tmp_final):
            raise RuntimeError("итоговый файл не прошёл проверку ffprobe")
        got = probe(tmp_final)["dur"]
        exp = ((e or dur) - s) if tc and (e or dur) else (dur if dur else None)
        if exp and abs(got - exp) > 1.5:
            notes.append(f"длительность {got:.1f} с вместо ожидаемых {exp:.1f} с")
        final.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(tmp_final), str(final))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return notes, info


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", help="xlsx / csv / txt / ссылка / поисковый запрос")
    ap.add_argument("--sheet", default="Монтажный лист")
    ap.add_argument("-o", "--out", help="папка выгрузки клипов (по умолчанию: clips/ рядом с исходным файлом)")
    ap.add_argument("--cache-dir", help="папка кэша сырых видео (по умолчанию: .cache внутри папки выгрузки)")
    ap.add_argument("--status", help="брать только строки с этим статусом (по умолчанию: пусто/найти/ошибка…)")
    ap.add_argument("--redo", action="store_true", help="игнорировать статус, обработать все строки")
    ap.add_argument("--priority", help="например A или AB"); ap.add_argument("--only", help="номера: 5,9-14")
    ap.add_argument("--limit", type=int)
    ap.add_argument("--tc", help="таймкод для одиночного режима"); ap.add_argument("--name", help="имя файла в одиночном режиме")
    ap.add_argument("--pick", action="store_true"); ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--results", type=int, default=5); ap.add_argument("--max-min", type=float, default=20)
    ap.add_argument("--allow-long", action="store_true", help="разрешить ролики длиннее лимита без таймкода")
    ap.add_argument("--height", type=int, default=1080)
    ap.add_argument("--cut-mode", choices=["auto", "full", "section"], default="auto",
                    help="full: скачать целиком и точно вырезать; section: скачать только отрезок (экономит трафик у длинных)")
    ap.add_argument("--full-max-min", type=float, default=30, help="auto: длиннее этого -> section")
    ap.add_argument("--margin", type=float, default=3, help="запас вокруг таймкода в режиме section, сек")
    ap.add_argument("--clip-len", type=float, default=15, help="длина клипа, если задана одна точка времени")
    ap.add_argument("--attempts", type=int, default=2); ap.add_argument("--sleep", type=float, default=0, help="пауза между скачиваниями, сек")
    ap.add_argument("--cookies-from-browser", dest="cfb"); ap.add_argument("--cookies", help="cookies.txt")
    ap.add_argument("--proxy", help="например socks5://127.0.0.1:1080")
    ap.add_argument("--clean-cache", action="store_true"); ap.add_argument("--verbose", action="store_true")
    args = ap.parse_args()

    if args.out:
        out = Path(args.out)
    else:
        src_p = Path(args.source)
        if src_p.is_file():
            out = src_p.resolve().parent / "clips"
        else:
            out = Path("clips").resolve()
    out.mkdir(parents=True, exist_ok=True)
    args.resolved_out = out
    args.resolved_cache = Path(args.cache_dir) if args.cache_dir else (out / ".cache")
    run_log = out / f"run_{dt.datetime.now():%Y%m%d_%H%M%S}.txt"
    lf = open(run_log, "w", encoding="utf-8")
    sys.stdout, sys.stderr = Tee(sys.__stdout__, lf), Tee(sys.__stderr__, lf)
    print(f"Лог этого запуска: {run_log}\nКоманда: {' '.join(sys.argv)}")
    print(f"Python {platform.python_version()} | {platform.platform()} | yt-dlp {yt_dlp.version.__version__}")
    for tool in ("ffmpeg", "ffprobe", "deno"):
        path = shutil.which(tool)
        ver = ""
        if path:
            try:
                ver = subprocess.run([tool, "-version" if tool != "deno" else "--version"], capture_output=True, text=True, timeout=10).stdout.splitlines()[0]
            except Exception:
                pass
        print(f"{tool}: {path or 'НЕ НАЙДЕН'} {ver}")
    if not shutil.which("ffmpeg") or not shutil.which("ffprobe"):
        sys.exit("ffmpeg/ffprobe не найдены в PATH.")
    if not shutil.which("deno"):
        print("ВНИМАНИЕ: не найден deno — без него YouTube часто отдаёт 403. winget install DenoLand.Deno (затем новое окно терминала)", file=sys.stderr)

    items, sync = load_items(args.source, args)
    only = parse_only(args.only) if args.only else None
    todo = [it for it in items if is_pending(it, args) and (not only or it["n"] in only)
            and (not args.priority or it["prio"] in args.priority) and it["query"] not in ("—", "-")]
    if args.limit:
        todo = todo[:args.limit]
    print(f"Позиций в файле: {len(items)}, к обработке: {len(todo)}")

    log_new = not (out / "log.csv").exists()
    logf = open(out / "log.csv", "a", newline="", encoding="utf-8-sig")
    log = csv.writer(logf)
    if log_new:
        log.writerow(["№", "блок", "момент", "запрос", "таймкод", "ссылка", "название", "длительность", "канал", "файл", "результат", "заметка"])
    stats = {"ok": 0, "skip": 0, "fail": []}

    for idx, it in enumerate(todo, 1):
        label = f"[{idx}/{len(todo)}] №{it['n']}"
        folder = out / slug(it["block"], 30) if it["block"] else out
        final = folder / f"{int(it['n'] or idx):02d}_{slug(it['title'] or it['query'])}.mp4"
        row = [it["n"], it["block"], it["title"], it["query"], it["tc"]]
        try:
            tc = parse_tc(it["tc"], args.clip_len)
        except ValueError as e:
            msg = f"ошибка: таймкод — {e}"
            print(f"{label} {msg}"); stats["fail"].append((it["n"], msg))
            log.writerow(row + ["", "", "", "", "", msg, ""]); logf.flush()
            if sync and not args.dry_run: sync.update(it, status=msg)
            continue
        if is_valid(final):
            print(f"{label} уже скачан: {final}")
            stats["skip"] += 1
            if sync and not args.dry_run: sync.update(it, status="скачано", file=str(final.resolve()))
            continue
        if final.exists():
            print(f"{label} файл повреждён, удаляю и качаю заново"); final.unlink()
        result = None
        try:
            if re.match(r"https?://", it["query"]):
                cand = {"url": it["query"], "title": "(прямая ссылка)", "dur": None, "chan": "", "live": False}
            else:
                cand = choose(search(it["query"], args), it, tc is not None, args)
            if not cand:
                raise SkipItem("нет подходящих кандидатов (см. «отсеяно»)")
            print(f"{label} {cand['title'][:70]} [{fmt_dur(cand['dur'])}] {cand['url']}")
            if args.dry_run:
                log.writerow(row + [cand["url"], cand["title"], fmt_dur(cand["dur"]), cand["chan"], "", "dry-run", ""]); logf.flush()
                continue
            for attempt in range(args.attempts):
                try:
                    notes, info = process(it, cand["url"], tc, args, out, final, label)
                    result = ("скачано", notes, info)
                    break
                except SkipItem:
                    raise
                except Exception as e:
                    if attempt + 1 >= args.attempts:
                        raise
                    wait = 5 * 2 ** attempt
                    print(f"{label} сбой ({str(e)[:120]}), повтор через {wait} с")
                    time.sleep(wait)
            status, notes, info = result
            print(f"{label} OK -> {final}" + (f"   ⚠ {'; '.join(notes)}" if notes else ""))
            stats["ok"] += 1
            title = info.get("title") or cand["title"]
            log.writerow(row + [cand["url"], title, fmt_dur(info.get("duration")), cand["chan"], str(final), "ok", "; ".join(notes)])
            if sync: sync.update(it, status="скачано", url=cand["url"], file=str(final.resolve()), note="; ".join(notes) or None)
        except Exception as e:
            msg = str(e).replace("\n", " ")[:300]
            hint = ""
            if re.search(r"403|Sign in|not a bot|Precondition|JavaScript", msg, re.I):
                hint = " | подсказка: yt-dlp -U / pip install -U \"yt-dlp[default]\", установить deno, --cookies-from-browser chrome, --proxy"
            print(f"{label} ошибка: {msg}{hint}", file=sys.stderr)
            if not isinstance(e, SkipItem):
                print(traceback.format_exc())
            stats["fail"].append((it["n"], msg))
            log.writerow(row + ["", "", "", "", "", "ошибка", msg])
            if sync and not args.dry_run: sync.update(it, status=f"ошибка: {msg[:120]}", note=hint.strip(" |") or None)
        logf.flush()
        if args.sleep:
            time.sleep(random.uniform(args.sleep, args.sleep * 2))

    logf.close()
    shutil.rmtree(out / "_tmp", ignore_errors=True)
    if args.clean_cache:
        c_dir = getattr(args, "resolved_cache", out / ".cache")
        shutil.rmtree(c_dir, ignore_errors=True)
    print(f"\nИтог: скачано {stats['ok']}, уже было {stats['skip']}, ошибок {len(stats['fail'])}")
    for n, m in stats["fail"]:
        print(f"  №{n}: {m}")
    print(f"Лог: {run_log}")


if __name__ == "__main__":
    main()
