from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Any, Dict
import openpyxl


@dataclass
class ProjectSettings:
    fps: float = 25.0
    width: int = 1920
    height: int = 1080
    handle_sec: float = 1.5
    min_speed: float = 0.8
    shortage_policy: List[str] = field(default_factory=lambda: ["slow", "gap", "flag"])
    default_duration_sec: float = 3.0
    voiceover_path: str = "voiceover.wav"
    script_path: str = "script.txt"
    whisper_model: str = "large-v3"
    language: str = "ru"
    fuzzy_threshold: float = 65.0

    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> "ProjectSettings":
        cfg = cls()
        if not data:
            return cfg

        for k, v in data.items():
            if v is None:
                continue
            k_clean = str(k).strip().lower().replace(" ", "_")
            if k_clean == "fps":
                cfg.fps = float(v)
            elif k_clean == "width":
                cfg.width = int(float(v))
            elif k_clean == "height":
                cfg.height = int(float(v))
            elif k_clean == "handle_sec":
                cfg.handle_sec = float(v)
            elif k_clean == "min_speed":
                cfg.min_speed = float(v)
            elif k_clean == "shortage_policy":
                if isinstance(v, list):
                    cfg.shortage_policy = [str(x).strip().lower() for x in v]
                else:
                    cfg.shortage_policy = [p.strip().lower() for p in str(v).split(",") if p.strip()]
            elif k_clean == "default_duration_sec":
                cfg.default_duration_sec = float(v)
            elif k_clean == "voiceover_path":
                cfg.voiceover_path = str(v).strip()
            elif k_clean == "script_path":
                cfg.script_path = str(v).strip()
            elif k_clean == "whisper_model":
                cfg.whisper_model = str(v).strip()
            elif k_clean == "language":
                cfg.language = str(v).strip()
            elif k_clean == "fuzzy_threshold":
                cfg.fuzzy_threshold = float(v)
        return cfg


def load_settings_from_workbook(wb: openpyxl.Workbook) -> ProjectSettings:
    """Loads settings from sheet 'Настройки' if present, otherwise returns defaults."""
    names = {s.strip().lower(): s for s in wb.sheetnames}
    sheet_name = names.get("настройки") or names.get("settings")
    if not sheet_name:
        return ProjectSettings()

    ws = wb[sheet_name]
    raw_data: Dict[str, Any] = {}
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 2:
            continue
        key, val = row[0], row[1]
        if key is not None and str(key).strip():
            raw_data[str(key).strip()] = val

    return ProjectSettings.from_dict(raw_data)


def load_settings(path: str | Path) -> ProjectSettings:
    """Loads settings from an Excel file."""
    p = Path(path)
    if not p.exists() or p.suffix.lower() != ".xlsx":
        return ProjectSettings()
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    try:
        return load_settings_from_workbook(wb)
    finally:
        wb.close()
