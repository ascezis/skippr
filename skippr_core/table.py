from pathlib import Path
import re
from typing import List, Tuple, Optional, Any, Dict
import openpyxl

from .models import RowItem
from .config import ProjectSettings, load_settings_from_workbook


NEW_COLUMNS = [
    "№",
    "Блок",
    "Фраза",
    "Вид",
    "Референс",
    "Поисковый запрос",
    "Таймкод (от-до)",
    "Длительность, с",
    "Смещение, с",
    "Приоритет",
    "Статус",
    "Заметки",
]

LEGACY_KIND_MAP = {
    "фильм": "клип",
    "музыка/мем": "клип",
    "мем": "клип",
    "тв/мем": "клип",
    "хроника/тв": "клип",
    "хроника": "клип",
    "тв": "клип",
    "скриншот": "картинка",
    "картинка": "картинка",
    "статья": "текст",
    "текст": "текст",
    "оригинал": "оригинал",
}


def _parse_legacy_duration(usage_text: Any, default_sec: float) -> float:
    if not usage_text:
        return default_sec
    s = str(usage_text).lower().replace(",", ".")
    # Look for patterns like "10–15 с", "3–4 с", "0.5 с"
    m = re.findall(r"(\d+(?:\.\d+)?)\s*(?:с|sec|s)\b", s)
    if m:
        try:
            # If multiple numbers (like 10-15s), pick the upper or first
            nums = [float(x) for x in m]
            return max(nums)
        except Exception:
            pass
    return default_sec


def load_table(
    path: str | Path, settings: Optional[ProjectSettings] = None
) -> Tuple[List[RowItem], ProjectSettings, bool]:
    """
    Loads cue sheet table from Excel.
    Returns (items, settings, is_legacy).
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"File not found: {p}")

    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    try:
        if settings is None:
            settings = load_settings_from_workbook(wb)

        # Find sheet
        names = {s.strip().lower(): s for s in wb.sheetnames}
        sheet_name = names.get("монтажный лист") or names.get("cue sheet") or wb.sheetnames[0]
        ws = wb[sheet_name]

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], settings, False

        header = [str(c or "").strip() for c in rows[0]]
        col_index: Dict[str, int] = {h: i for i, h in enumerate(header) if h}

        # Check if new format
        is_new_format = "Фраза" in col_index and "Вид" in col_index
        items: List[RowItem] = []

        if is_new_format:
            for r_idx, row in enumerate(rows[1:], start=2):
                if not any(row):
                    continue

                def val(name: str) -> Any:
                    idx = col_index.get(name)
                    return row[idx] if idx is not None and idx < len(row) else None

                n_raw = val("№")
                if n_raw is None or str(n_raw).strip() == "":
                    continue

                try:
                    n = int(float(n_raw))
                except Exception:
                    continue

                dur_raw = val("Длительность, с")
                try:
                    duration = float(dur_raw) if dur_raw is not None else settings.default_duration_sec
                except Exception:
                    duration = settings.default_duration_sec

                offset_raw = val("Смещение, с")
                try:
                    offset = float(offset_raw) if offset_raw is not None else 0.0
                except Exception:
                    offset = 0.0

                item = RowItem(
                    n=n,
                    block=str(val("Блок") or "").strip(),
                    phrase=str(val("Фраза") or "").strip(),
                    kind=str(val("Вид") or "клип").strip().lower(),
                    reference=str(val("Референс") or "").strip(),
                    query=str(val("Поисковый запрос") or "").strip(),
                    tc=str(val("Таймкод (от-до)") or "").strip() or None,
                    duration=duration,
                    offset=offset,
                    priority=str(val("Приоритет") or "B").strip().upper(),
                    status=str(val("Статус") or "найти").strip().lower(),
                    notes=str(val("Заметки") or "").strip(),
                    file=str(val("Файл") or "").strip() or None,
                    row_idx=r_idx,
                )
                items.append(item)
            return items, settings, False

        # Legacy format fallback (montazhny_list_opensorsa.xlsx)
        for r_idx, row in enumerate(rows[1:], start=2):
            if not any(row):
                continue

            def leg_val(name: str) -> Any:
                idx = col_index.get(name)
                return row[idx] if idx is not None and idx < len(row) else None

            n_raw = leg_val("№")
            if n_raw is None:
                continue
            try:
                n = int(float(n_raw))
            except Exception:
                continue

            moment = str(leg_val("Момент сценария") or "").strip()
            raw_kind = str(leg_val("Тип") or "").strip().lower()
            kind = LEGACY_KIND_MAP.get(raw_kind, "клип")
            usage = leg_val("Как использовать")
            duration = _parse_legacy_duration(usage, settings.default_duration_sec)

            item = RowItem(
                n=n,
                block=str(leg_val("Блок") or "").strip(),
                phrase=moment,
                kind=kind,
                reference=str(leg_val("Референс") or "").strip(),
                query=str(leg_val("Поисковый запрос") or "").strip(),
                tc=str(leg_val("Таймкод (от-до)") or "").strip() or None,
                duration=duration,
                offset=0.0,
                priority=str(leg_val("Приоритет") or "B").strip().upper(),
                status=str(leg_val("Статус") or "найти").strip().lower(),
                notes=str(usage or "").strip(),
                row_idx=r_idx,
            )
            items.append(item)

        return items, settings, True
    finally:
        wb.close()
