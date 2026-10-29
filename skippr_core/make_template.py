#!/usr/bin/env python3
"""
make_template.py — Generates a new Excel cue-sheet template for skippr pipeline.
Includes dropdown validations, auto-filters, 'Настройки' and 'Легенда' sheets.
"""
import argparse
from pathlib import Path
import sys
import openpyxl

for _s in (sys.stdout, sys.stderr):
    if _s and hasattr(_s, "reconfigure"):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from skippr_core.table import NEW_COLUMNS

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ROW_FONT = Font(name="Calibri", size=11)
BORDER_THIN = Side(border_style="thin", color="D9D9D9")
BORDER_BOX = Border(left=BORDER_THIN, right=BORDER_THIN, top=BORDER_THIN, bottom=BORDER_THIN)


def create_template(output_path: str | Path = "montazhny_list_template.xlsx") -> Path:
    out = Path(output_path)
    wb = openpyxl.Workbook()

    # ─────────────────────────────────────────────────────────────
    # Sheet 1: Монтажный лист
    # ─────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Монтажный лист"
    ws.freeze_panes = "A2"

    for col_idx, col_name in enumerate(NEW_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Example rows (optional guidelines)
    sample_rows = [
        (1, "Вступление", "Это история о том как создавался Linux", "клип", "Хроника 90-х Торвальдс", "Linus Torvalds 1991 interview", "0:15-0:18", 3.0, 0.0, "A", "найти", "Динамичная перебивка"),
        (2, "Вступление", "миллионы строк открытого кода", "текст", "Скриншот репозитория Git", "", None, 2.5, 0.0, "B", "найти", "Скриншот или плашка"),
        (3, "Вступление", "но корпорации были против", "картинка", "Логотипы компаний", "", None, 3.0, 0.0, "C", "найти", "Заглушка под графику"),
    ]

    for r_idx, row_data in enumerate(sample_rows, start=2):
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = ROW_FONT
            cell.border = BORDER_BOX
            if c_idx in (1, 8, 9):
                cell.alignment = Alignment(horizontal="center")

    # Data Validations (Dropdowns)
    # Kind (D)
    dv_kind = DataValidation(type="list", formula1='"клип,картинка,оригинал,текст"', allow_blank=True)
    ws.add_data_validation(dv_kind)
    dv_kind.add("D2:D500")

    # Priority (J)
    dv_prio = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True)
    ws.add_data_validation(dv_prio)
    dv_prio.add("J2:J500")

    # Status (K)
    dv_status = DataValidation(type="list", formula1='"найти,скачано,ошибка,пропуск"', allow_blank=True)
    ws.add_data_validation(dv_status)
    dv_status.add("K2:K500")

    # Enable AutoFilter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(NEW_COLUMNS))}500"

    # Column widths
    col_widths = {
        1: 6,   # №
        2: 18,  # Блок
        3: 35,  # Фраза
        4: 12,  # Вид
        5: 25,  # Референс
        6: 35,  # Поисковый запрос
        7: 18,  # Таймкод
        8: 15,  # Длительность, с
        9: 14,  # Смещение, с
        10: 12, # Приоритет
        11: 12, # Статус
        12: 30, # Заметки
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ─────────────────────────────────────────────────────────────
    # Sheet 2: Настройки
    # ─────────────────────────────────────────────────────────────
    ws_settings = wb.create_sheet(title="Настройки")
    settings_headers = ["Параметр", "Значение", "Описание"]
    for col_idx, h in enumerate(settings_headers, start=1):
        cell = ws_settings.cell(row=1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    settings_data = [
        ("fps", 25, "Частота кадров проекта DaVinci Resolve"),
        ("width", 1920, "Ширина кадра таймлайна"),
        ("height", 1080, "Высота кадра таймлайна"),
        ("handle_sec", 1.5, "Запас по краям клипа для подгонки монтажером (сек)"),
        ("min_speed", 0.8, "Минимальное замедление при нехватке хронометража"),
        ("shortage_policy", "slow,gap,flag", "Политика при нехватке длины (через запятую: slow, gap, flag)"),
        ("default_duration_sec", 3.0, "Длительность клипа по умолчанию, если не указана (сек)"),
        ("voiceover_path", "voiceover.wav", "Путь к аудиофайлу озвучки (wav или mp3)"),
        ("script_path", "script.txt", "Путь к текстовому файлу сценария"),
        ("whisper_model", "large-v3", "Модель faster-whisper (tiny/base/small/medium/large-v3)"),
        ("language", "ru", "Язык озвучки для распознавания"),
    ]

    for r_idx, (param, val, desc) in enumerate(settings_data, start=2):
        c1 = ws_settings.cell(row=r_idx, column=1, value=param)
        c2 = ws_settings.cell(row=r_idx, column=2, value=val)
        c3 = ws_settings.cell(row=r_idx, column=3, value=desc)
        c1.font = Font(name="Calibri", size=11, bold=True)
        c2.font = ROW_FONT
        c3.font = ROW_FONT
        c1.border = c2.border = c3.border = BORDER_BOX

    ws_settings.column_dimensions["A"].width = 24
    ws_settings.column_dimensions["B"].width = 20
    ws_settings.column_dimensions["C"].width = 55

    # ─────────────────────────────────────────────────────────────
    # Sheet 3: Легенда
    # ─────────────────────────────────────────────────────────────
    ws_legend = wb.create_sheet(title="Легенда")
    legend_headers = ["Колонка", "Тип", "Правила и назначение"]
    for col_idx, h in enumerate(legend_headers, start=1):
        cell = ws_legend.cell(row=1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    legend_data = [
        ("№", "Число", "Уникальный порядковый номер вставки."),
        ("Блок", "Текст", "Имя смысловой части ролика. В него группируются файлы: clips/<Блок>/."),
        ("Фраза", "Текст", "Дословная фраза из сценария. Служит якорем для синхронизации со звуком озвучки."),
        ("Вид", "Список", "клип — скачиваемое видео. картинка/оригинал/текст — заглушки на таймлайне."),
        ("Референс", "Текст", "Описание сцены или отсылки для человека / редактора."),
        ("Поисковый запрос", "Текст", "Запрос на YouTube (на английском) или прямая ссылка (https://...)."),
        ("Таймкод (от-до)", "Текст", "Диапазон вырезки: 1:23-1:40, 85-110, 0:45+15s. Необязательно."),
        ("Длительность, с", "Число", "Целевая длительность вставки на таймлайне в секундах."),
        ("Смещение, с", "Число", "Сдвиг начала клипа относительно произнесения фразы (по умолчанию 0)."),
        ("Приоритет", "A / B / C", "A/B — основная дорожка V1. C — опциональная дорожка V2."),
        ("Статус", "Список", "найти — скачать. скачано — готово. ошибка — сбой. пропуск — игнорировать."),
        ("Заметки", "Текст", "Любые заметки монтажеру, скриптом не читаются."),
    ]

    for r_idx, (col_n, t_type, desc) in enumerate(legend_data, start=2):
        c1 = ws_legend.cell(row=r_idx, column=1, value=col_n)
        c2 = ws_legend.cell(row=r_idx, column=2, value=t_type)
        c3 = ws_legend.cell(row=r_idx, column=3, value=desc)
        c1.font = Font(name="Calibri", size=11, bold=True)
        c2.font = ROW_FONT
        c3.font = ROW_FONT
        c1.border = c2.border = c3.border = BORDER_BOX

    ws_legend.column_dimensions["A"].width = 20
    ws_legend.column_dimensions["B"].width = 15
    ws_legend.column_dimensions["C"].width = 75

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main():
    parser = argparse.ArgumentParser(description="Generate empty skippr cue sheet template.")
    parser.add_argument("-o", "--output", default="montazhny_list_template.xlsx", help="Output xlsx path")
    args = parser.parse_args()
    p = create_template(args.output)
    print(f"Шаблон создан: {p.resolve()}")


if __name__ == "__main__":
    main()
