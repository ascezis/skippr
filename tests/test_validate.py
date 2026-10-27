import pytest
import openpyxl
from pathlib import Path
from make_template import create_template
from validate import validate_cue_sheet
from skippr_core.table import load_table


def test_make_template(tmp_path):
    tpl_path = tmp_path / "montazhny_list_template.xlsx"
    res = create_template(tpl_path)
    assert res.exists()

    wb = openpyxl.load_workbook(res)
    assert "Монтажный лист" in wb.sheetnames
    assert "Настройки" in wb.sheetnames
    assert "Легенда" in wb.sheetnames

    ws = wb["Монтажный лист"]
    headers = [cell.value for cell in ws[1]]
    assert "Фраза" in headers
    assert "Вид" in headers
    assert "Длительность, с" in headers
    wb.close()


def test_validate_success(tmp_path):
    table_path = tmp_path / "test_table.xlsx"
    create_template(table_path)

    # Script contains phrases
    script_path = tmp_path / "script.txt"
    script_path.write_text(
        "Это история о том как создавался Linux, где миллионы строк открытого кода, но корпорации были против.",
        encoding="utf-8",
    )

    items, errors = validate_cue_sheet(table_path, script_path)
    crit_errors = [e for e in errors if not e.is_warning]
    assert len(crit_errors) == 0
    assert len(items) == 3


def test_validate_catches_duplicate_n(tmp_path):
    table_path = tmp_path / "test_table.xlsx"
    create_template(table_path)

    wb = openpyxl.load_workbook(table_path)
    ws = wb["Монтажный лист"]
    # Change row 3's № from 2 to 1 (duplicate)
    ws.cell(row=3, column=1, value=1)
    wb.save(table_path)
    wb.close()

    script_path = tmp_path / "script.txt"
    script_path.write_text("Это история о том как создавался Linux миллионы строк открытого кода но корпорации были против", encoding="utf-8")

    _, errors = validate_cue_sheet(table_path, script_path)
    dup_errors = [e for e in errors if e.field == "№" and "Дублирующийся" in e.message]
    assert len(dup_errors) > 0


def test_validate_catches_empty_duration(tmp_path):
    table_path = tmp_path / "test_table.xlsx"
    create_template(table_path)

    wb = openpyxl.load_workbook(table_path)
    ws = wb["Монтажный лист"]
    # Clear duration on row 2 (column 8)
    ws.cell(row=2, column=8, value=None)
    wb.save(table_path)
    wb.close()

    script_path = tmp_path / "script.txt"
    script_path.write_text("Это история о том как создавался Linux миллионы строк открытого кода но корпорации были против", encoding="utf-8")

    # In load_table, if duration is None, it defaults to settings.default_duration_sec,
    # unless invalid value like <= 0 is passed
    ws_wb = openpyxl.load_workbook(table_path)
    ws2 = ws_wb["Монтажный лист"]
    ws2.cell(row=2, column=8, value=-5.0)
    ws_wb.save(table_path)
    ws_wb.close()

    _, errors = validate_cue_sheet(table_path, script_path)
    dur_errors = [e for e in errors if e.field == "Длительность"]
    assert len(dur_errors) > 0


def test_validate_catches_phrase_not_in_script(tmp_path):
    table_path = tmp_path / "test_table.xlsx"
    create_template(table_path)

    script_path = tmp_path / "script.txt"
    # Script does not mention "миллионы строк открытого кода"
    script_path.write_text("Совсем другой текст сценария про космос и ракеты.", encoding="utf-8")

    _, errors = validate_cue_sheet(table_path, script_path)
    phrase_errors = [e for e in errors if e.field == "Фраза" and "не найдена" in e.message]
    assert len(phrase_errors) > 0


def test_validate_catches_invalid_kind(tmp_path):
    table_path = tmp_path / "test_table.xlsx"
    create_template(table_path)

    wb = openpyxl.load_workbook(table_path)
    ws = wb["Монтажный лист"]
    ws.cell(row=2, column=4, value="неизвестный_вид_файла")
    wb.save(table_path)
    wb.close()

    script_path = tmp_path / "script.txt"
    script_path.write_text("Это история о том как создавался Linux миллионы строк открытого кода но корпорации были против", encoding="utf-8")

    _, errors = validate_cue_sheet(table_path, script_path)
    kind_errors = [e for e in errors if e.field == "Вид" and "Недопустимое значение" in e.message]
    assert len(kind_errors) > 0
