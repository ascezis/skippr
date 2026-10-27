import pytest
import subprocess
import json
from pathlib import Path
import openpyxl

from make_template import create_template
from skippr_core.table import NEW_COLUMNS
from skippr import step_all


@pytest.fixture
def test_project_dir(tmp_path):
    proj = tmp_path / "test_project"
    proj.mkdir(parents=True, exist_ok=True)

    # 1. Generate cue sheet table with 3 NVIDIA rows
    table_path = proj / "montazhny_list.xlsx"
    create_template(table_path)

    wb = openpyxl.load_workbook(table_path)
    ws = wb["Монтажный лист"]

    # Clear sample rows and insert NVIDIA test rows
    for r in range(2, 20):
        for c in range(1, 15):
            ws.cell(row=r, column=c, value=None)

    nvidia_rows = [
        (1, "NVIDIA", "Вопрос зрительницы и реакция Линуса", "клип", "Aalto Talk with Linus Torvalds", "https://youtu.be/mock1", None, 3.0, 0.0, "A", "найти", "Ключевая сцена"),
        (2, "NVIDIA", "Инди корпорация Nvidia", "клип", "Дженсен Хуанг", "https://youtu.be/mock2", None, 2.5, 0.0, "B", "найти", "Шутка"),
        (3, "NVIDIA", "Премия Millennium Technology Prize", "картинка", "Церемония вручения 2012", "", None, 2.0, 0.0, "C", "найти", "Фото/Заглушка"),
    ]

    for r_idx, row_vals in enumerate(nvidia_rows, start=2):
        for c_idx, val in enumerate(row_vals, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    # Update settings sheet
    ws_set = wb["Настройки"]
    ws_set.cell(row=9, column=2, value="voiceover.wav")
    ws_set.cell(row=10, column=2, value="script.txt")

    wb.save(table_path)
    wb.close()

    # 2. Create script.txt
    script_file = proj / "script.txt"
    script_file.write_text(
        "В этом блоке разбирается вопрос зрительницы и реакция Линуса на тему драйверов. "
        "Инди корпорация Nvidia стала отдельным интернет-мемом, пока наконец не была вручена "
        "премия Millennium Technology Prize за развитие открытого программного обеспечения.",
        encoding="utf-8",
    )

    # 3. Create synthetic voiceover.wav (sine tone 15s)
    vo_file = proj / "voiceover.wav"
    cmd = [
        "ffmpeg", "-y", "-hide_banner", "-loglevel", "error",
        "-f", "lavfi", "-i", "sine=frequency=440:duration=15",
        "-c:a", "pcm_s16le", str(vo_file),
    ]
    subprocess.run(cmd, check=True)

    # 4. Create pre-computed transcript.json cache
    cache_dir = proj / "cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    transcript_file = cache_dir / "transcript.json"

    words = [
        {"word": "В", "start": 0.0, "end": 0.2, "probability": 1.0},
        {"word": "этом", "start": 0.2, "end": 0.5, "probability": 1.0},
        {"word": "блоке", "start": 0.5, "end": 0.9, "probability": 1.0},
        {"word": "разбирается", "start": 0.9, "end": 1.5, "probability": 1.0},
        {"word": "вопрос", "start": 1.5, "end": 1.9, "probability": 1.0},
        {"word": "зрительницы", "start": 1.9, "end": 2.5, "probability": 1.0},
        {"word": "и", "start": 2.5, "end": 2.6, "probability": 1.0},
        {"word": "реакция", "start": 2.6, "end": 3.0, "probability": 1.0},
        {"word": "Линуса", "start": 3.0, "end": 3.6, "probability": 1.0},
        {"word": "на", "start": 3.6, "end": 3.8, "probability": 1.0},
        {"word": "тему", "start": 3.8, "end": 4.1, "probability": 1.0},
        {"word": "драйверов", "start": 4.1, "end": 4.7, "probability": 1.0},
        {"word": "Инди", "start": 5.0, "end": 5.4, "probability": 1.0},
        {"word": "корпорация", "start": 5.4, "end": 6.1, "probability": 1.0},
        {"word": "Nvidia", "start": 6.1, "end": 6.7, "probability": 1.0},
        {"word": "стала", "start": 6.7, "end": 7.0, "probability": 1.0},
        {"word": "отдельным", "start": 7.0, "end": 7.5, "probability": 1.0},
        {"word": "премия", "start": 8.0, "end": 8.5, "probability": 1.0},
        {"word": "Millennium", "start": 8.5, "end": 9.2, "probability": 1.0},
        {"word": "Technology", "start": 9.2, "end": 9.8, "probability": 1.0},
        {"word": "Prize", "start": 9.8, "end": 10.4, "probability": 1.0},
    ]

    transcript_data = {
        "duration": 15.0,
        "language": "ru",
        "segments": [{"start": 0.0, "end": 15.0, "text": script_file.read_text(encoding="utf-8")}],
        "words": words,
    }
    transcript_file.write_text(json.dumps(transcript_data, ensure_ascii=False, indent=2), encoding="utf-8")

    return proj


def test_pipeline_dry_run_criterion_6(test_project_dir):
    """
    Acceptance Criterion 6:
    Полный прогон python skippr.py all --dry-run на тестовых данных завершается без исключений и пишет report.md.
    """
    table = test_project_dir / "montazhny_list.xlsx"

    # Run complete pipeline in dry-run mode
    step_all(table_path=str(table), out_dir=str(test_project_dir), dry_run=True)

    # Verify generated outputs
    phrases_file = test_project_dir / "phrases.json"
    plan_file = test_project_dir / "plan.json"
    xml_file = test_project_dir / "timeline.fcpxml"
    report_file = test_project_dir / "report.md"

    assert phrases_file.exists(), "phrases.json must be generated"
    assert plan_file.exists(), "plan.json must be generated"
    assert xml_file.exists(), "timeline.fcpxml must be generated"
    assert report_file.exists(), "report.md must be generated"

    # Verify phrases.json contents
    phrases = json.loads(phrases_file.read_text(encoding="utf-8"))
    assert len(phrases) == 3
    assert phrases[0]["method"] == "exact"
    assert phrases[1]["method"] == "exact"
    assert phrases[2]["method"] == "exact"

    # Verify plan.json contents
    plan = json.loads(plan_file.read_text(encoding="utf-8"))
    assert len(plan) == 3
    assert plan[0]["track"] == "V1"
    assert plan[1]["track"] == "V1"
    assert plan[2]["track"] == "V2"  # Priority C

    # Verify report.md
    report_text = report_file.read_text(encoding="utf-8")
    assert "# Отчёт конвейера skippr" in report_text
    assert "NVIDIA" in report_text


def test_pipeline_with_edl(test_project_dir):
    table = test_project_dir / "montazhny_list.xlsx"
    step_all(table_path=str(table), out_dir=str(test_project_dir), dry_run=True, edl=True)

    edl_file = test_project_dir / "timeline.edl"
    assert edl_file.exists()
    content = edl_file.read_text(encoding="utf-8")
    assert "TITLE: Skippr Timeline" in content
    assert "FCM: NON-DROP FRAME" in content

